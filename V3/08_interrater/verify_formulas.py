"""
Verify that the workbook's Excel formulas compute the statistics they claim
===========================================================================

    python 08_interrater/verify_formulas.py        # exits non-zero on mismatch

The Agreement sheet of MLATE_V3_interrater_MASTER.xlsx is built entirely from
live Excel formulas, so that the three raters' returned columns can simply be
pasted in. That is convenient and completely unverifiable by inspection: an
Excel formula that silently evaluates to zero looks exactly like one that is
correct until someone checks the arithmetic.

This script checks it. It fills a throwaway copy of the workbook with simulated
ratings of known structure, evaluates every formula with a pure-Python Excel
engine, and compares the result against an independent NumPy implementation of
Cohen's kappa, quadratic-weighted kappa and Fleiss' kappa written from the
definitions. Every reported statistic must agree to 1e-6.

This is not a formality. The first version of the builder expressed the
per-row agreement indicators as `--(A2=B2)`, the usual Excel idiom for coercing
a boolean to 1/0. Under evaluation every one of them returned zero, which
silently zeroed every percent-agreement figure and drove all six Cohen's kappas
negative, while the weighted kappas - which are built from `(A2-B2)^2` and never
touch the coercion - stayed correct. A reader glancing at the workbook would
have seen plausible weighted kappas beside impossible unweighted ones. The
indicators are now written as `IF(A2=B2,1,0)`, which is unambiguous in both
Excel and the checker.

Requires the `formulas` package (`pip install formulas`).
"""

from __future__ import annotations

import re
import shutil
import sys
import tempfile
import warnings
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mlate import config as cfg

warnings.filterwarnings("ignore")

MASTER = cfg.RESULTS_DIR / "08_interrater" / "MLATE_V3_interrater_MASTER.xlsx"
CATS = {"Printability": [0, 1, 2, 3], "Cell Response": [1, 2, 3, 4, 5]}
RATING_COL = {"Printability": 14, "Cell Response": 15}
REF_COL = {"Printability": 6, "Cell Response": 7}
BLANK_ROWS = [3, 47, 128]          # zero-based; exercises pairwise exclusion
TOL = 1e-6


# ─────────────────────────────────────────────────────────────────────────────
# Reference implementation, written from the definitions
# ─────────────────────────────────────────────────────────────────────────────
def cohen_kappa(a, b, cats, weighted=False):
    n, k = len(a), len(cats)
    pa = np.array([(a == c).sum() for c in cats], dtype=float) / n
    pb = np.array([(b == c).sum() for c in cats], dtype=float) / n
    if not weighted:
        po, pe = (a == b).mean(), float(pa @ pb)
        return (po - pe) / (1 - pe)
    W = np.array([[(i - j) ** 2 / (k - 1) ** 2 for j in cats] for i in cats])
    do = np.mean((a - b) ** 2) / (k - 1) ** 2
    return 1 - do / float(pa @ W @ pb)


def fleiss_kappa(mat, cats):
    N, n = mat.shape
    counts = np.array([[(row == c).sum() for c in cats] for row in mat],
                      dtype=float)
    Pi = ((counts ** 2).sum(axis=1) - n) / (n * (n - 1))
    pbar = float(Pi.mean())
    pj = counts.sum(axis=0) / (N * n)
    pbar_e = float(pj @ pj)
    return pbar, pbar_e, (pbar - pbar_e) / (1 - pbar_e)


def simulate(path: Path, seed: int = 7):
    """Write simulated ratings into a copy of the master and return the truth."""
    wb = load_workbook(path)
    ref_ws = wb["Reference_Labels"]
    # Column A also holds a trailing explanatory note, so count only the
    # cells that actually carry a Row_ID.
    n = sum(1 for r in range(2, ref_ws.max_row + 1)
            if re.fullmatch(r"R\d+",
                            str(ref_ws.cell(row=r, column=1).value or "")))
    rows = list(range(2, n + 2))

    rng = np.random.default_rng(seed)
    truth, sim = {}, {}
    for target, cats in CATS.items():
        t = np.array([ref_ws.cell(row=r, column=REF_COL[target]).value
                      for r in rows], dtype=float)
        truth[target] = t
        arms = []
        # Three raters of deliberately different reliability, so a bug that
        # collapses the pairwise statistics onto one value is visible.
        for keep in (0.82, 0.78, 0.62):
            vals = []
            for v in t:
                u = rng.random()
                if u < keep:
                    x = v
                elif u < keep + (1 - keep) * 0.75:
                    x = v + rng.choice([-1, 1])
                else:
                    x = v + rng.choice([-2, 2])
                vals.append(float(np.clip(x, cats[0], cats[-1])))
            arms.append(np.array(vals))
        sim[target] = arms

    for rater in range(3):
        ws = wb[f"Rater_{rater + 1}"]
        for target in CATS:
            col = RATING_COL[target]
            for i, r in enumerate(rows):
                blank = rater == 1 and i in BLANK_ROWS
                ws.cell(row=r, column=col).value = (
                    None if blank else int(sim[target][rater][i]))
    wb.save(path)

    for target in CATS:          # mirror the blanks in the expected values
        sim[target][1] = sim[target][1].copy()
        sim[target][1][BLANK_ROWS] = np.nan
    return truth, sim


def expected(truth, sim):
    out = {}
    for target, cats in CATS.items():
        R, ref = sim[target], truth[target]
        complete = ~np.isnan(R[0]) & ~np.isnan(R[1]) & ~np.isnan(R[2])
        C = [r[complete] for r in R]
        s = {"Formulations rated by all three raters": float(complete.sum()),
             "All three raters identical":
                 float(np.mean((C[0] == C[1]) & (C[1] == C[2])))}
        ex, adj, ck, wk = [], [], [], []
        for tag, (i, j) in zip(("rater 1 vs 2", "rater 1 vs 3", "rater 2 vs 3"),
                               ((0, 1), (0, 2), (1, 2))):
            s[f"Exact agreement, {tag}"] = float(np.mean(C[i] == C[j]))
            ex.append(s[f"Exact agreement, {tag}"])
            adj.append(float(np.mean(np.abs(C[i] - C[j]) <= 1)))
            s[f"Cohen's kappa, {tag}"] = cohen_kappa(C[i], C[j], cats)
            ck.append(s[f"Cohen's kappa, {tag}"])
            s[f"Weighted kappa, {tag}"] = cohen_kappa(C[i], C[j], cats, True)
            wk.append(s[f"Weighted kappa, {tag}"])
        s["Exact agreement, mean over the three pairs"] = float(np.mean(ex))
        s["Agreement within one class, mean over pairs"] = float(np.mean(adj))
        s["Cohen's kappa, mean over the three pairs"] = float(np.mean(ck))
        s["Weighted kappa, mean over the three pairs"] = float(np.mean(wk))
        pbar, pbar_e, fk = fleiss_kappa(np.column_stack(C), cats)
        s["Mean observed agreement, P-bar"] = pbar
        s["Chance agreement, P-bar-e"] = pbar_e
        s["Fleiss' kappa"] = fk
        pr = np.array([(ref[complete] == c).mean() for c in cats])
        for k in (1, 2, 3):
            a = R[k - 1]
            m = ~np.isnan(a) & ~np.isnan(ref)
            po = float(np.mean(a[m] == ref[m]))
            pa = np.array([(C[k - 1] == c).mean() for c in cats])
            pe = float(pa @ pr)
            s[f"Rater {k}: exact match with the dataset label"] = po
            s[f"Rater {k}: Cohen's kappa against the dataset label"] = \
                (po - pe) / (1 - pe)
        out[target] = s
    return out


def evaluate(path: Path) -> dict:
    """Column A label -> column B value, per target section, from the engine."""
    import formulas
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()

    def scalar(v):
        try:
            return v.value[0, 0]
        except Exception:
            return v

    cells = {}
    for key, v in sol.items():
        ku = key.upper()
        if "AGREEMENT'!A" in ku or "AGREEMENT'!B" in ku:
            col = "A" if "AGREEMENT'!A" in ku else "B"
            num = "".join(ch for ch in ku.rsplit("!", 1)[1] if ch.isdigit())
            if num:
                cells.setdefault(int(num), {})[col] = scalar(v)

    out, target = {}, None
    for r in sorted(cells):
        label, value = cells[r].get("A"), cells[r].get("B")
        label = "" if label is None else str(label).strip()
        for t in CATS:
            if label.startswith(f"{t}   (") or label.startswith(f"{t}  ("):
                target = t
                out[t] = {}
        if target and label and value is not None and not isinstance(value, str):
            out[target][label] = value
    return out


def main() -> None:
    if not MASTER.exists():
        raise SystemExit(f"build the workbook first: {MASTER}")
    tmp = Path(tempfile.mkdtemp()) / "verify.xlsx"
    shutil.copy(MASTER, tmp)
    truth, sim = simulate(tmp)
    exp, got = expected(truth, sim), evaluate(tmp)

    print(f"{'statistic':52s} {'workbook':>10s} {'expected':>10s}   ")
    print("-" * 84)
    failures = checked = 0
    for target in CATS:
        print(f"\n== {target} ==")
        for label, want in exp[target].items():
            have = got.get(target, {}).get(label)
            if have is None:
                print(f"{label[:52]:52s} {'MISSING':>10s} {want:10.4f}   FAIL")
                failures += 1
                continue
            checked += 1
            ok = abs(float(have) - want) < TOL
            print(f"{label[:52]:52s} {float(have):10.4f} {want:10.4f}   "
                  f"{'ok' if ok else 'FAIL'}")
            failures += 0 if ok else 1
    print("-" * 84)
    print(f"{checked - failures} of {checked} statistics agree to {TOL:g}")
    shutil.rmtree(tmp.parent, ignore_errors=True)
    if failures:
        raise SystemExit(f"{failures} formula(s) disagree with the reference "
                         f"implementation")
    print("all workbook formulas verified")


if __name__ == "__main__":
    main()
